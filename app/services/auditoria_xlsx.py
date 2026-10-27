from io import BytesIO
from datetime import datetime
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter


# Estilos reutilizáveis
HEADER_FONT = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
HEADER_FILL = PatternFill(start_color='1F3A5F', end_color='1F3A5F', fill_type='solid')
SUBHEADER_FILL = PatternFill(start_color='2E5C8A', end_color='2E5C8A', fill_type='solid')
TOTAL_FILL = PatternFill(start_color='DCE6F0', end_color='DCE6F0', fill_type='solid')
ALT_ROW_FILL = PatternFill(start_color='F2F5F9', end_color='F2F5F9', fill_type='solid')
WHITE_FILL = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

TITLE_FONT = Font(name='Calibri', bold=True, size=16, color='1F3A5F')
SUBTITLE_FONT = Font(name='Calibri', bold=True, size=12, color='2E5C8A')
NORMAL_FONT = Font(name='Calibri', size=10)
BOLD_FONT = Font(name='Calibri', bold=True, size=10)
CARD_VALUE_FONT = Font(name='Calibri', bold=True, size=14, color='1F3A5F')
CARD_LABEL_FONT = Font(name='Calibri', size=9, color='555555')

THIN_BORDER = Border(
    left=Side(style='thin', color='CCCCCC'),
    right=Side(style='thin', color='CCCCCC'),
    top=Side(style='thin', color='CCCCCC'),
    bottom=Side(style='thin', color='CCCCCC'),
)

HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
CENTER_ALIGN = Alignment(horizontal='center', vertical='center')
RIGHT_ALIGN = Alignment(horizontal='right', vertical='center')
LEFT_ALIGN = Alignment(horizontal='left', vertical='center')


def _valor_brl(valor) -> str:
    try:
        v = float(valor or 0)
    except (TypeError, ValueError):
        v = 0
    return f"R$ {v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _status_label(status: str) -> str:
    labels = {
        "CONFORME": "Conforme",
        "ATENCAO": "Atenção",
        "CRITICO": "Crítico",
    }
    return labels.get(status, status)


def _gravidade_label(gravidade: str) -> str:
    labels = {"ALTA": "Alta", "MEDIA": "Média", "BAIXA": "Baixa"}
    return labels.get(gravidade, gravidade)


def _status_pendencia_label(status: str) -> str:
    labels = {"FALTANTE": "Faltante", "INCONSISTENTE": "Inconsistente"}
    return labels.get(status, status)


def _tipo_doc_label(tipo: str) -> str:
    mapa = {
        "pgdas_d": "PGDAS-D",
        "iss": "ISS",
        "das": "DAS",
        "folha_pagamento": "Folha de Pagamento",
        "defis": "DEFIS",
        "dirf": "DIRF",
        "efd_reinf": "EFD-REINF",
        "dap": "DAP",
    }
    return mapa.get(tipo.lower(), tipo.upper())


def _apply_header_style(ws, row: int, max_col: int):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER


def _apply_subheader_style(ws, row: int, max_col: int):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = SUBHEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER


def _apply_total_style(ws, row: int, max_col: int):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = BOLD_FONT
        cell.fill = TOTAL_FILL
        cell.border = THIN_BORDER


def _apply_cell_style(ws, row: int, col: int, bold: bool = False, align: Alignment = LEFT_ALIGN, number_format: str = None):
    cell = ws.cell(row=row, column=col)
    cell.font = BOLD_FONT if bold else NORMAL_FONT
    cell.alignment = align
    cell.border = THIN_BORDER
    if number_format:
        cell.number_format = number_format
    if row % 2 == 0:
        cell.fill = ALT_ROW_FILL
    else:
        cell.fill = WHITE_FILL


def _auto_width(ws, min_width: int = 10, max_width: int = 50):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        adjusted = min(max(max_len + 2, min_width), max_width)
        ws.column_dimensions[col_letter].width = adjusted


def gerar_xlsx_auditoria(
    periodo: str,
    empresas: list[dict],
) -> bytes:
    """
    Gera planilha XLSX com duas abas:
    1. RESUMO - KPIs consolidados + resumo por empresa
    2. DETALHAMENTO - Todas as pendências de todas as empresas
    """
    wb = Workbook()

    # ==================== ABA 1: RESUMO ====================
    ws1 = wb.active
    ws1.title = "RESUMO"

    row = 1
    # Título
    ws1.merge_cells('A1:H1')
    cell = ws1.cell(row=1, column=1, value="2RS Notas - Relatório de Auditoria")
    cell.font = TITLE_FONT
    cell.alignment = Alignment(horizontal='center', vertical='center')
    row += 1

    ws1.merge_cells(f'A{row}:H{row}')
    cell = ws1.cell(row=row, column=1, value=f"Competência: {periodo}")
    cell.font = SUBTITLE_FONT
    cell.alignment = Alignment(horizontal='center', vertical='center')
    row += 1

    ws1.merge_cells(f'A{row}:H{row}')
    cell = ws1.cell(row=row, column=1, value=f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}")
    cell.font = NORMAL_FONT
    cell.alignment = Alignment(horizontal='center', vertical='center')
    row += 2

    # KPIs Cards (linha de valores)
    kpi_headers = ["Total Empresas", "Conformes", "Em Atenção", "Críticas", "Total Documentos Faltantes"]
    total_empresas = len(empresas)
    conformes = sum(1 for e in empresas if e.get("status_geral") == "CONFORME")
    criticas = sum(1 for e in empresas if e.get("status_geral") == "CRITICO")
    atencao = total_empresas - conformes - criticas
    total_faltantes = sum(
        (e.get("documentos_esperados", 0) - e.get("documentos_encontrados", 0))
        for e in empresas
    )
    kpi_values = [total_empresas, conformes, atencao, criticas, total_faltantes]

    for col, (header, value) in enumerate(zip(kpi_headers, kpi_values), 1):
        ws1.cell(row=row, column=col, value=header).font = CARD_LABEL_FONT
        ws1.cell(row=row, column=col).alignment = CENTER_ALIGN
        cell = ws1.cell(row=row + 1, column=col, value=value)
        cell.font = CARD_VALUE_FONT
        cell.alignment = CENTER_ALIGN
    row += 3

    # Tabela: Resumo por Empresa
    ws1.merge_cells(f'A{row}:H{row}')
    cell = ws1.cell(row=row, column=1, value="Resumo por Empresa")
    cell.font = SUBTITLE_FONT
    row += 1

    empresa_headers = [
        "Razão Social", "CNPJ", "Status Geral", "Conformidade (%)",
        "Fator R (%)", "Alíquota Efetiva (%)", "Anexo",
        "Docs Encontrados / Esperados", "Pendências Críticas"
    ]
    for col, header in enumerate(empresa_headers, 1):
        ws1.cell(row=row, column=col, value=header)
    _apply_header_style(ws1, row, len(empresa_headers))
    row += 1

    for emp in empresas:
        ws1.cell(row=row, column=1, value=emp.get("razao_social", "-"))
        ws1.cell(row=row, column=2, value=emp.get("cnpj", "-"))
        ws1.cell(row=row, column=3, value=_status_label(emp.get("status_geral", "-")))
        ws1.cell(row=row, column=4, value=round(emp.get("conformidade_pct", 0) or 0, 2))
        ws1.cell(row=row, column=4).number_format = '0.00"%"'
        ws1.cell(row=row, column=5, value=round(emp.get("fator_r", 0) or 0, 2))
        ws1.cell(row=row, column=5).number_format = '0.00"%"'
        ws1.cell(row=row, column=6, value=round(emp.get("aliquota_efetiva", 0) or 0, 2))
        ws1.cell(row=row, column=6).number_format = '0.00"%"'
        ws1.cell(row=row, column=7, value=emp.get("classificacao_anexo", "-") or "-")
        ws1.cell(row=row, column=8, value=f"{emp.get('documentos_encontrados', 0)} / {emp.get('documentos_esperados', 0)}")
        ws1.cell(row=row, column=9, value=emp.get("pendencias_criticas_count", 0))

        for col in range(1, len(empresa_headers) + 1):
            _apply_cell_style(ws1, row, col)
        row += 1

    # Linha TOTAL
    ws1.cell(row=row, column=1, value="TOTAL GERAL")
    ws1.cell(row=row, column=4, value=round(sum(e.get("conformidade_pct", 0) or 0 for e in empresas) / max(total_empresas, 1), 2))
    ws1.cell(row=row, column=4).number_format = '0.00"%"'
    ws1.cell(row=row, column=8, value=f"{sum(e.get('documentos_encontrados', 0) for e in empresas)} / {sum(e.get('documentos_esperados', 0) for e in empresas)}")
    ws1.cell(row=row, column=9, value=sum(e.get("pendencias_criticas_count", 0) for e in empresas))
    _apply_total_style(ws1, row, len(empresa_headers))
    row += 2

    # Ajustar larguras da aba RESUMO
    widths_resumo = [40, 18, 15, 15, 12, 15, 10, 25, 18]
    for i, w in enumerate(widths_resumo, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    # ==================== ABA 2: DETALHAMENTO ====================
    ws2 = wb.create_sheet("DETALHAMENTO")

    row = 1
    ws2.merge_cells('A1:L1')
    cell = ws2.cell(row=1, column=1, value="2RS Notas - Detalhamento de Pendências")
    cell.font = TITLE_FONT
    cell.alignment = Alignment(horizontal='center', vertical='center')
    row += 1

    ws2.merge_cells(f'A{row}:L{row}')
    cell = ws2.cell(row=row, column=1, value=f"Competência: {periodo}")
    cell.font = SUBTITLE_FONT
    cell.alignment = Alignment(horizontal='center', vertical='center')
    row += 1

    ws2.merge_cells(f'A{row}:L{row}')
    cell = ws2.cell(row=row, column=1, value=f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}")
    cell.font = NORMAL_FONT
    cell.alignment = Alignment(horizontal='center', vertical='center')
    row += 2

    pend_headers = [
        "Empresa", "CNPJ", "Status Empresa", "Tipo Documento",
        "Gravidade", "Status Pendência", "Observação",
        "CRF (%)", "CRT", "Base Cálculo", "Receita Declarada", "Conformidade Empresa (%)"
    ]
    for col, header in enumerate(pend_headers, 1):
        ws2.cell(row=row, column=col, value=header)
    _apply_header_style(ws2, row, len(pend_headers))
    row += 1

    for emp in empresas:
        pendencias = emp.get("pendencias", [])
        if not pendencias:
            # Linha indicando sem pendências
            ws2.cell(row=row, column=1, value=emp.get("razao_social", "-"))
            ws2.cell(row=row, column=2, value=emp.get("cnpj", "-"))
            ws2.cell(row=row, column=3, value=_status_label(emp.get("status_geral", "-")))
            ws2.merge_cells(start_row=row, start_column=4, end_row=row, end_column=len(pend_headers))
            cell = ws2.cell(row=row, column=4, value="Sem pendências registradas")
            cell.font = Font(name='Calibri', size=10, italic=True, color='888888')
            cell.alignment = Alignment(horizontal='center', vertical='center')
            for col in range(1, len(pend_headers) + 1):
                _apply_cell_style(ws2, row, col)
            row += 1
            continue

        for p in pendencias:
            ws2.cell(row=row, column=1, value=emp.get("razao_social", "-"))
            ws2.cell(row=row, column=2, value=emp.get("cnpj", "-"))
            ws2.cell(row=row, column=3, value=_status_label(emp.get("status_geral", "-")))
            ws2.cell(row=row, column=4, value=_tipo_doc_label(p.get("tipo_documento", "-")))
            ws2.cell(row=row, column=5, value=_gravidade_label(p.get("gravidade", "-")))
            ws2.cell(row=row, column=6, value=_status_pendencia_label(p.get("status", "-")))
            ws2.cell(row=row, column=7, value=p.get("observacao", "-") or "-")
            ws2.cell(row=row, column=8, value=round(p.get("crf", 0) or 0, 2))
            ws2.cell(row=row, column=8).number_format = '0.00"%"'
            ws2.cell(row=row, column=9, value=p.get("crt", "-") or "-")
            ws2.cell(row=row, column=10, value=p.get("base_calculo", 0) or 0)
            ws2.cell(row=row, column=10).number_format = '#,##0.00'
            ws2.cell(row=row, column=11, value=p.get("receita_declarada", 0) or 0)
            ws2.cell(row=row, column=11).number_format = '#,##0.00'
            ws2.cell(row=row, column=12, value=round(emp.get("conformidade_pct", 0) or 0, 2))
            ws2.cell(row=row, column=12).number_format = '0.00"%"'

            for col in range(1, len(pend_headers) + 1):
                _apply_cell_style(ws2, row, col)
            row += 1

    # Ajustar larguras da aba DETALHAMENTO
    widths_detalhe = [40, 18, 15, 20, 12, 15, 40, 12, 10, 15, 15, 18]
    for i, w in enumerate(widths_detalhe, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # Congelar painéis (cabeçalho fixo)
    ws1.freeze_panes = 'A6'
    ws2.freeze_panes = 'A5'

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()